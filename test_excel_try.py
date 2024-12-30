import openpyxl

from openpyxl import Workbook, load_workbook

# load, excel file
wb = load_workbook('C:\\Users\\PRIVATE_ILANA\\PHYTHON_HOW_TO\\pythonProject\\excel_for_python_exp.xlsx')

# read the work sheet of the excel file
ws = wb.active
print(f"The name of the Work Sheet is: {ws}")
wb_sheet = wb['My bank acount excel file']

# change the name for the excel work sheet
#wb_sheet.title = "My bank acount excel file"

# change the color of the excel work sheet
wb_sheet.sheet_properties.tabColor = 'FF0000'

# save all changes back into a file
wb.save('C:\\Users\\PRIVATE_ILANA\\PHYTHON_HOW_TO\\pythonProject\\excel_for_python_exp.xlsx')

# read value from specific cell
transaction_fee_B3 = ws['B3'].value
print(f"The value I just read is: {transaction_fee_B3}")

# write value into specific cell
ws['B6'].value = 25000
# save all changes back into a file
wb.save('C:\\Users\\PRIVATE_ILANA\\PHYTHON_HOW_TO\\pythonProject\\excel_for_python_exp.xlsx')

# merge cells







